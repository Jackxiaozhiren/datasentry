"""Step 19 修复引擎测试（12.5/15 章 MVP 子集 + ADR-020，M6 验收）。"""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest

from datasentry_core.connectors import (
    CsvConnector,
    DataSourceSpec,
    DataSourceType,
    default_registry,
)
from datasentry_core.detectors import DetectionContext, DetectorRegistry
from datasentry_core.detectors.initial import register_default_detectors
from datasentry_core.detectors.runner import ScanRunner
from datasentry_core.models.enums import (
    RepairOperation,
    RepairRunStatus,
)
from datasentry_core.models.repair import RepairProposal
from datasentry_core.repair import RepairEngine
from datasentry_core.repair.engine import changed_cells, rows_equal

RULES = (
    "name,status,price,event_date\n"
    + "\n".join(f"user{i},active,{i * 10},2024-01-01" for i in range(20))
    + "\n"
    + " user9 ,Active,n/a,2024-02-30\n"
    + "  user8  ,active,300,2024-13-01\n"
)


def _ctx(tmp_path: Path, csv_text: str) -> tuple[DetectionContext, Path, DetectorRegistry]:
    p = tmp_path / "data.csv"
    p.write_text(csv_text, encoding="utf-8")
    spec = DataSourceSpec(source_type=DataSourceType.CSV, path=p, options={"dataset_id": "t"})
    handle = CsvConnector().open(spec)
    context = DetectionContext(
        dataset_id="t",
        table_name=None,
        columns=handle.schema().column_names,
        handle=handle,
    )
    registry = DetectorRegistry()
    register_default_detectors(registry)
    return context, p, registry


def _close(context: DetectionContext) -> None:
    context.handle.close()


def _run_scan(context: DetectionContext, registry: DetectorRegistry) -> list:
    runner = ScanRunner(registry)
    _, _, issues = runner.run_scan(context, None)
    return issues


def _issue_with_detector(issues: list, detector_id: str):
    """家族化 issue 里挑出含指定原始检测器（与 issue_type 同名）的。"""
    return next(i for i in issues if detector_id in i.detector_ids)


class TestPropose:
    def test_proposes_trim_for_whitespace(self, tmp_path: Path) -> None:
        ctx, _, registry = _ctx(tmp_path, RULES)
        try:
            issues = _run_scan(ctx, registry)
            whitespace = _issue_with_detector(issues, "leading_or_trailing_whitespace")
            proposal = RepairEngine().propose(whitespace, ctx)
            assert proposal is not None
            assert proposal.operation == RepairOperation.TRIM_WHITESPACE
            assert proposal.issue_type == "leading_or_trailing_whitespace"
            assert proposal.target_columns == ["name"]
            assert proposal.estimated_rows_changed == 2
        finally:
            _close(ctx)

    def test_proposes_token_replacement(self, tmp_path: Path) -> None:
        ctx, _, registry = _ctx(tmp_path, RULES)
        try:
            issues = _run_scan(ctx, registry)
            token = _issue_with_detector(issues, "suspicious_missing_token")
            proposal = RepairEngine().propose(token, ctx)
            assert proposal is not None
            assert proposal.operation == RepairOperation.REPLACE_MISSING_TOKEN
        finally:
            _close(ctx)

    def test_no_proposal_for_unmapped_issue(self, tmp_path: Path) -> None:
        ctx, _, registry = _ctx(tmp_path, RULES)
        try:
            issues = _run_scan(ctx, registry)
            unmatched = _issue_with_detector(issues, "uniqueness_violation")
            proposal = RepairEngine().propose(unmatched, ctx)
            assert proposal is None
        finally:
            _close(ctx)


class TestPreview:
    def test_preview_reports_deltas_and_rule_failures(self, tmp_path: Path) -> None:
        ctx, _, registry = _ctx(tmp_path, RULES)
        try:
            issues = _run_scan(ctx, registry)
            whitespace = _issue_with_detector(issues, "leading_or_trailing_whitespace")
            engine = RepairEngine()
            proposal = engine.propose(whitespace, ctx)
            assert proposal is not None
            preview = engine.preview(proposal, ctx, registry)
            assert preview.rows_changed == 2
            assert preview.rule_failures_before["leading_or_trailing_whitespace"] > 0
            assert preview.rule_failures_after["leading_or_trailing_whitespace"] == 0
            assert (
                preview.rule_failures_after["leading_or_trailing_whitespace"]
                < (preview.rule_failures_before["leading_or_trailing_whitespace"])
            )
            assert preview.changed_examples
            example = preview.changed_examples[0]
            assert example.column == "name"
            assert str(example.before) != str(example.after)
        finally:
            _close(ctx)


class TestApplyRollback:
    def test_apply_creates_repaired_copy(self, tmp_path: Path) -> None:
        ctx, source_path, registry = _ctx(tmp_path, RULES)
        workspace = tmp_path / "ws"
        try:
            issues = _run_scan(ctx, registry)
            whitespace = _issue_with_detector(issues, "leading_or_trailing_whitespace")
            engine = RepairEngine()
            proposal = engine.propose(whitespace, ctx)
            assert proposal is not None
            source_before = hashlib.sha256(source_path.read_bytes()).hexdigest()
            run = engine.apply(proposal, ctx, workspace)
            assert run.status == RepairRunStatus.APPLIED
            assert run.fingerprint_before != run.fingerprint_after
            output = workspace / ".datasentry" / "repairs" / f"{run.id}.csv"
            assert output.exists()
            # D3-01：不变量 3——apply 永不覆盖源文件（源 sha256 前后一致）
            assert hashlib.sha256(source_path.read_bytes()).hexdigest() == source_before
            artifact = Path(run.rollback_artifact or "")
            assert artifact.exists()
            # 修复副本上重跑检测器：0 候选
            after_spec = DataSourceSpec(
                source_type=DataSourceType.CSV, path=output, options={"dataset_id": "t"}
            )
            after_handle = CsvConnector().open(after_spec)
            try:
                after_ctx = DetectionContext(
                    dataset_id="t",
                    table_name=None,
                    columns=after_handle.schema().column_names,
                    handle=after_handle,
                )
                detector = registry.get("leading_or_trailing_whitespace")
                assert detector.detect(after_ctx) == []
            finally:
                after_handle.close()
        finally:
            _close(ctx)

    def test_rollback_restores_before_state(self, tmp_path: Path) -> None:
        ctx, _, registry = _ctx(tmp_path, RULES)
        workspace = tmp_path / "ws"
        try:
            issues = _run_scan(ctx, registry)
            whitespace = _issue_with_detector(issues, "leading_or_trailing_whitespace")
            engine = RepairEngine()
            proposal = engine.propose(whitespace, ctx)
            assert proposal is not None
            run = engine.apply(proposal, ctx, workspace)
            rolled = engine.rollback(run, workspace)
            assert rolled.status == RepairRunStatus.ROLLED_BACK
            rolled_path = workspace / ".datasentry" / "repairs" / f"{run.id}.rolled_back.csv"
            assert rolled_path.exists()
            # 回滚副本指纹 = before 指纹（M6：回滚后状态一致）
            before = ctx.handle.fingerprint()
            rolled_spec = DataSourceSpec(
                source_type=DataSourceType.CSV,
                path=rolled_path,
                options={"dataset_id": "t"},
            )
            rolled_handle = CsvConnector().open(rolled_spec)
            try:
                after_fp = rolled_handle.fingerprint()
                assert after_fp.file_sha256 == before.file_sha256
            finally:
                rolled_handle.close()
        finally:
            _close(ctx)


class TestClip:
    def test_clip_proposal_from_outlier_bounds(self, tmp_path: Path) -> None:
        rows = [f"{i}" for i in range(100)] + ["5000", "-5000"]
        ctx, _, registry = _ctx(tmp_path, "value\n" + "\n".join(rows) + "\n")
        try:
            issues = _run_scan(ctx, registry)
            outlier = _issue_with_detector(issues, "iqr_outlier")
            engine = RepairEngine()
            proposal = engine.propose(outlier, ctx)
            assert proposal is not None
            assert proposal.operation == RepairOperation.CLIP_VALUE
            assert "lower" in proposal.parameters and "upper" in proposal.parameters
            preview = engine.preview(proposal, ctx, registry)
            assert preview.rule_failures_after["iqr_outlier"] == 0
        finally:
            _close(ctx)


ZIPS = "id,zip\n1, 01234 \n2, 09876 \n3, 00700 \n"


def _write_pair(tmp_path: Path, before: str, after: str) -> tuple[Path, Path]:
    b = tmp_path / "before.csv"
    a = tmp_path / "after.csv"
    b.write_text(before, encoding="utf-8")
    a.write_text(after, encoding="utf-8")
    return b, a


class TestTableDiffComparesRawText:
    """D2-06（r7）：diff 比较的是工件里的原文，不是连接器推断后的值。

    pyarrow 默认把 ` 01234 ` 推断成 int64 `1234`，before/after 两侧同时失真，于是一次改光
    全表的 TRIM 在 diff 里报成"什么都没改"——可审计环节对已发生的变更说谎。既有
    `test_repair_diff_cli` 用的是 name 列（字符串列，推断不改值），所以 1350 例全绿也测不到。
    """

    def test_padded_numeric_column_reports_every_changed_row(self, tmp_path: Path) -> None:
        b, a = _write_pair(tmp_path, ZIPS, 'id,zip\n1,"01234"\n2,"09876"\n3,"00700"\n')
        columns, before_rows, after_rows, changed = RepairEngine.table_diff(
            b, a, DataSourceType.CSV
        )
        assert columns == ["id", "zip"]
        assert changed == [0, 1, 2]
        assert before_rows[0] == ["1", " 01234 "]
        assert after_rows[0] == ["1", "01234"]

    def test_value_formatting_difference_is_visible(self, tmp_path: Path) -> None:
        """`1.50` 与 `1.5` 是同一台数字、两段不同的字节——审计要看的是后者。"""
        b, a = _write_pair(tmp_path, "x\n1.50\n", "x\n1.5\n")
        assert RepairEngine.table_diff(b, a, DataSourceType.CSV)[3] == [0]

    def test_row_count_difference_is_not_truncated_away(self, tmp_path: Path) -> None:
        """一侧少一行时必须报成变更，不能像 zip(strict=False) 那样静默截断。"""
        b, a = _write_pair(tmp_path, "x\n1\n2\n3\n", "x\n1\n2\n")
        _columns, before_rows, after_rows, changed = RepairEngine.table_diff(
            b, a, DataSourceType.CSV
        )
        assert changed == [2]
        assert len(before_rows) == 3 and len(after_rows) == 2

    def test_identical_artifacts_report_no_changes(self, tmp_path: Path) -> None:
        """对照组：前两例的"非空"不能来自把一切判成变更。"""
        b, a = _write_pair(tmp_path, ZIPS, ZIPS)
        assert RepairEngine.table_diff(b, a, DataSourceType.CSV)[3] == []

    def test_a_real_trim_repair_shows_up_in_its_own_diff(self, tmp_path: Path) -> None:
        """端到端：真实 apply 产出的两份工件，diff 必须报出 3/3 行变更。"""
        ctx, _, registry = _ctx(tmp_path, ZIPS)
        workspace = tmp_path / "ws"
        try:
            issues = _run_scan(ctx, registry)
            whitespace = _issue_with_detector(issues, "leading_or_trailing_whitespace")
            engine = RepairEngine()
            proposal = engine.propose(whitespace, ctx)
            assert proposal is not None
            run = engine.apply(proposal, ctx, workspace)
            copy_path = workspace / ".datasentry" / "repairs" / f"{run.id}.csv"
            _cols, before_rows, after_rows, changed = RepairEngine.table_diff(
                Path(run.rollback_artifact or ""), copy_path, DataSourceType.CSV
            )
            assert changed == [0, 1, 2], "diff 对一次改了全部数据行的 TRIM 报成无变更 —— 即 D2-06"
            assert before_rows[0][1] == " 01234 " and after_rows[0][1] == "01234"
        finally:
            _close(ctx)


class TestArtefactReadFidelity:
    """复核 B-1…B-4：`table_diff` 用第二套读取规则读工件，而工件是连接器写下的。

    扫描侧会嗅探分隔符（`csv.Sniffer`，`,;\\t|`）、取 `worksheets[0]`、按 `\\n` 切 JSONL；
    审计侧一概不会。于是同一份字节被读了两遍并得到两个答案——diff 就把没发生的变更报成
    发生了（`None -> 'Alice'`），或者干脆崩掉。本类的每一例都直接喂两份工件，判的是
    "审计侧读出的，是不是文件里真有的东西"。
    """

    def _pair(self, tmp_path: Path, name: str, before: str, after: str) -> tuple[Path, Path]:
        b = tmp_path / f"before{name}"
        a = tmp_path / f"after{name}"
        b.write_text(before, encoding="utf-8")
        a.write_text(after, encoding="utf-8")
        return b, a

    def test_a_tab_delimited_pair_is_read_as_tab_delimited(self, tmp_path: Path) -> None:
        b, a = self._pair(
            tmp_path, ".tsv", "id\tname\n1\t  Alice \n2\tBob\n", "id\tname\n1\tAlice\n2\tBob\n"
        )
        columns, before_rows, after_rows, changed = RepairEngine.table_diff(
            b, a, DataSourceType.CSV
        )
        assert columns == ["id", "name"], "the reader collapsed the header into one column"
        assert before_rows[0] == ["1", "  Alice "], before_rows
        assert after_rows[0] == ["1", "Alice"], after_rows
        assert changed == [0], f"fabricated changes: {changed}"

    def test_a_semicolon_delimited_pair_is_read_as_semicolon_delimited(
        self, tmp_path: Path
    ) -> None:
        b, a = self._pair(
            tmp_path, ".csv", "id;name\n1;  Alice \n2;Bob\n", "id;name\n1;Alice\n2;Bob\n"
        )
        columns, before_rows, _after, changed = RepairEngine.table_diff(b, a, DataSourceType.CSV)
        assert columns == ["id", "name"]
        assert before_rows[0] == ["1", "  Alice "], before_rows
        assert changed == [0], f"fabricated changes: {changed}"

    def test_duplicate_header_names_still_line_up_by_position(self, tmp_path: Path) -> None:
        """连接器把重名列改写成 `id`/`id_1`，快照里两段都叫 `id`——按名字投影会丢一整列。"""
        b, a = self._pair(
            tmp_path, ".csv", "id,id\n1,  Alice \n2,Bob\n", "id,id_1\n1,Alice\n2,Bob\n"
        )
        columns, before_rows, after_rows, changed = RepairEngine.table_diff(
            b, a, DataSourceType.CSV
        )
        assert columns == ["id", "id_1"]
        assert before_rows[0] == ["1", "  Alice "], f"lost a column: {before_rows}"
        assert after_rows[0] == ["1", "Alice"]
        assert changed == [0], f"fabricated changes: {changed}"

    def test_the_xlsx_snapshot_uses_the_sheet_the_connector_scans(self, tmp_path: Path) -> None:
        from openpyxl import Workbook

        def workbook(path: Path, name: str, active_notes: bool) -> Path:
            wb = Workbook()
            data = wb.active
            data.title = "Data"
            data.append(["id", "name"])
            data.append([1, name])
            notes = wb.create_sheet("Notes")
            notes.append(["x"])
            notes.append(["y"])
            if active_notes:
                wb.active = wb["Notes"]
            wb.save(path)
            return path

        b = workbook(tmp_path / "before.xlsx", "  Alice ", active_notes=True)
        a = workbook(tmp_path / "after.xlsx", "Alice", active_notes=True)
        columns, before_rows, after_rows, changed = RepairEngine.table_diff(
            b, a, DataSourceType.XLSX
        )
        assert columns == ["id", "name"], f"read the Notes sheet instead: {columns}"
        assert before_rows[0] == [1, "  Alice "], before_rows
        assert after_rows[0] == [1, "Alice"]
        assert changed == [0]

    def test_a_jsonl_record_containing_u2028_still_diffs(self, tmp_path: Path) -> None:
        """`str.splitlines()` 在 U+2028/U+2029/U+0085 上也断行，JSONL 的记录分隔符只有 `\\n`。"""
        import json

        before = "\n".join(
            [
                json.dumps({"a": "keep", "b": "  x "}, ensure_ascii=False),
                json.dumps({"a": "plain\u2028text", "b": "y"}, ensure_ascii=False),
            ]
        )
        after = "\n".join(
            [
                json.dumps({"a": "keep", "b": "x"}, ensure_ascii=False),
                json.dumps({"a": "plain\u2028text", "b": "y"}, ensure_ascii=False),
            ]
        )
        b, a = self._pair(tmp_path, ".jsonl", before + "\n", after + "\n")
        columns, before_rows, after_rows, changed = RepairEngine.table_diff(
            b, a, DataSourceType.JSONL
        )
        assert columns == ["a", "b"]
        assert before_rows[0] == ["keep", "  x "], before_rows
        assert after_rows[0] == ["keep", "x"], after_rows
        assert changed == [0], changed


class TestSidesShareOneSchema:
    """复核 F1/F2/F3/F4（追加 3h 的独立复核）：两份工件必须在**同一套列名**下对齐。

    上一轮把投影从"按名字"改成"按位置"，解决了连接器改名 `id`/`id_1` 丢列的问题，但列名仍然
    只取自 after 一侧，而每一侧的列名又来自 pyarrow 对**那一份文件**的推断。于是三份东西各自
    为政：before 独有的列被整列抹掉（说"什么都没改"），JSONL 的键顺序决定列顺序（说"改了"），
    `from_pylist` 只看第一条记录定型（后面的键静默消失），重名表头经 `dict(zip(...))` 塌成一列
    （后半张表的数据没了）。审计要回答"这一格变了吗"，前提是先有"同一格"。
    """

    def _parquet_pair(
        self, tmp_path: Path, before: dict[str, list[object]], after: dict[str, list[object]]
    ) -> tuple[Path, Path]:
        import pyarrow as pa
        import pyarrow.parquet as pq

        b = tmp_path / "before.parquet"
        a = tmp_path / "after.parquet"
        pq.write_table(pa.table(before), b)
        pq.write_table(pa.table(after), a)
        return b, a

    def test_a_column_present_only_in_the_snapshot_is_not_hidden(self, tmp_path: Path) -> None:
        b, a = self._parquet_pair(
            tmp_path,
            {"id": [1, 2], "name": ["x", "y"], "note": ["old", "chg"]},
            {"id": [1, 2], "name": ["x", "y"]},
        )
        columns, before_rows, after_rows, changed = RepairEngine.table_diff(
            b, a, DataSourceType.PARQUET
        )
        assert "note" in columns, f"a whole column vanished from the evidence: {columns}"
        assert before_rows[0] == [1, "x", "old"], before_rows
        assert after_rows[0] == [1, "x", None], after_rows
        assert changed == [0, 1], "a column vanished and the audit said nothing changed"

    def test_a_column_present_only_in_the_copy_is_reported_not_invented(
        self, tmp_path: Path
    ) -> None:
        b, a = self._parquet_pair(tmp_path, {"id": [1, 2]}, {"id": [1, 2], "extra": ["x", "y"]})
        columns, before_rows, after_rows, changed = RepairEngine.table_diff(
            b, a, DataSourceType.PARQUET
        )
        assert columns == ["id", "extra"]
        assert before_rows == [[1, None], [2, None]], before_rows
        assert after_rows == [[1, "x"], [2, "y"]], after_rows
        assert changed == [0, 1]

    def test_jsonl_records_with_the_same_keys_in_a_different_order_are_equal(
        self, tmp_path: Path
    ) -> None:
        """JSON 对象的键顺序不是数据。按位置投影两侧各自的推断顺序，会把同一条记录报成改过。"""
        b, a = (tmp_path / "before.jsonl", tmp_path / "after.jsonl")
        b.write_text('{"v":1,"w":2}\n', encoding="utf-8")
        a.write_text('{"w":2,"v":1}\n', encoding="utf-8")
        columns, before_rows, after_rows, changed = RepairEngine.table_diff(
            b, a, DataSourceType.JSONL
        )
        assert columns == ["v", "w"], columns
        assert before_rows == after_rows, (before_rows, after_rows)
        assert changed == [], f"key order fabricated a change: {changed}"

    def test_a_key_that_appears_only_after_the_first_record_survives_the_read(
        self, tmp_path: Path
    ) -> None:
        """`pa.Table.from_pylist` 只按第一条记录定型：第二条才出现的键会整列消失。"""
        b, a = (tmp_path / "before.jsonl", tmp_path / "after.jsonl")
        b.write_text('{"name":"  Alice "}\n{"name":"Bob","id":2}\n', encoding="utf-8")
        a.write_text('{"name":"Alice","id":1}\n{"name":"Bob","id":2}\n', encoding="utf-8")
        columns, before_rows, _after, changed = RepairEngine.table_diff(b, a, DataSourceType.JSONL)
        assert columns == ["name", "id"], columns
        assert before_rows[1] == ["Bob", 2], f"lost a later record's key: {before_rows}"
        assert changed == [0]

    def test_a_later_jsonl_record_that_omits_an_earlier_key_stays_in_its_column(
        self, tmp_path: Path
    ) -> None:
        """JSONL 允许某条记录少一个键。按"这条记录自己的值顺序"读会把后面的值挪到错的列上。"""
        b, a = (tmp_path / "before.jsonl", tmp_path / "after.jsonl")
        b.write_text('{"a":1,"b":2}\n{"b":3}\n', encoding="utf-8")
        a.write_text('{"a":1,"b":"x"}\n{"b":3}\n', encoding="utf-8")
        columns, before_rows, after_rows, changed = RepairEngine.table_diff(
            b, a, DataSourceType.JSONL
        )
        assert columns == ["a", "b"], columns
        assert before_rows[1] == [None, 3], f"a value slid into the wrong column: {before_rows}"
        assert after_rows[1] == [None, 3], after_rows
        assert changed == [0], changed

    def test_two_identical_jsonl_files_with_mixed_value_types_do_not_crash(
        self, tmp_path: Path
    ) -> None:
        path = tmp_path / "both.jsonl"
        path.write_text('{"v":1}\n{"v":"a"}\n', encoding="utf-8")
        columns, before_rows, after_rows, changed = RepairEngine.table_diff(
            path, path, DataSourceType.JSONL
        )
        assert columns == ["v"]
        assert before_rows == after_rows
        assert changed == [], f"identical bytes reported a change: {changed}"

    def test_an_xlsx_sheet_with_a_numeric_header_row_still_diffs(self, tmp_path: Path) -> None:
        from openpyxl import Workbook

        def workbook(path: Path, name: str) -> Path:
            wb = Workbook()
            ws = wb.active
            ws.append([1, 2])
            ws.append([name, 2])
            wb.save(path)
            return path

        b = workbook(tmp_path / "before.xlsx", "  Alice ")
        a = workbook(tmp_path / "after.xlsx", "Alice")
        columns, before_rows, _after, changed = RepairEngine.table_diff(b, a, DataSourceType.XLSX)
        assert columns == ["1", "2"], columns
        assert before_rows[0] == ["  Alice ", 2], before_rows
        assert changed == [0]

    def test_an_xlsx_sheet_with_duplicate_headers_keeps_both_columns(self, tmp_path: Path) -> None:
        """`dict(zip(cols,row))` 在 `cols=['id','id']` 时只留最后一个键：整张表左半边的数据没了。"""
        from openpyxl import Workbook

        def workbook(path: Path, left: str, right: str) -> Path:
            wb = Workbook()
            ws = wb.active
            ws.append(["id", "id"])
            ws.append([left, right])
            wb.save(path)
            return path

        b = workbook(tmp_path / "before.xlsx", "  Alice ", "Bob")
        a = workbook(tmp_path / "after.xlsx", "Alice", "Bob")
        columns, before_rows, after_rows, changed = RepairEngine.table_diff(
            b, a, DataSourceType.XLSX
        )
        assert columns == ["id", "id_1"], columns
        assert before_rows[0] == ["  Alice ", "Bob"], f"collapsed a column: {before_rows}"
        assert after_rows[0] == ["Alice", "Bob"]
        assert changed == [0]


class TestReadSideMatchesWriteSide:
    """复核 F1–F4/F6/F8（第二批独立复核）：读侧的命名、编码与谓词必须和写侧是同一套。

    上一轮把两侧统一到"按列名对齐"，但列名是**读侧自己造的**：空白表头格读成 `""` 而连接器写成
    `col_1`，去重造的 `a_1` 又会撞上文件里本来就有的 `a_1`。按名字对齐于是把一列裂成两列、把一整列
    读错一位——B-2 的"按名字丢列"从我自己的修法里复活。同一族的还有：读 CSV 只传了 delimiter 没传
    编码（非 UTF-8 源必然抛错），NaN 只修顶层，`preview` 没跟着加 guard，被拒的修复仍会先建目录。
    """

    def test_a_blank_xlsx_header_cell_is_named_the_way_the_connector_names_it(
        self, tmp_path: Path
    ) -> None:
        from openpyxl import Workbook

        def workbook(path: Path, header: list[object], rows: list[list[object]]) -> Path:
            wb = Workbook()
            ws = wb.active
            ws.append(header)
            for row in rows:
                ws.append(row)
            wb.save(path)
            return path

        b = workbook(tmp_path / "before.xlsx", ["id", None], [[1, "  Alice "], [2, "Bob"]])
        a = workbook(tmp_path / "after.xlsx", ["id", "col_1"], [[1, "Alice"], [2, "Bob"]])
        columns, before_rows, after_rows, changed = RepairEngine.table_diff(
            b, a, DataSourceType.XLSX
        )
        assert columns == ["id", "col_1"], f"one column split into two: {columns}"
        assert before_rows == [[1, "  Alice "], [2, "Bob"]], before_rows
        assert after_rows == [[1, "Alice"], [2, "Bob"]], after_rows
        assert changed == [0], f"every row reported as rewritten: {changed}"

    def test_a_non_utf8_snapshot_pair_is_decoded_per_side_instead_of_crashing(
        self, tmp_path: Path
    ) -> None:
        b = tmp_path / "before.csv"
        a = tmp_path / "after.csv"
        b.write_bytes(b"name,city\n  Alice ,Z\xfcrich\nBob,M\xfcnchen\n")
        a.write_bytes(b'"name","city"\n"Alice","Z\xc3\xbcrich"\n"Bob","M\xc3\xbcnchen"\n')
        columns, before_rows, after_rows, changed = RepairEngine.table_diff(
            b, a, DataSourceType.CSV
        )
        assert columns == ["name", "city"], columns
        assert before_rows[0] == ["  Alice ", "Zürich"], before_rows
        assert after_rows[0] == ["Alice", "Zürich"], after_rows
        assert changed == [0]

    def test_a_nan_inside_a_list_or_struct_is_not_a_change(self, tmp_path: Path) -> None:
        import math

        import pyarrow as pa
        import pyarrow.parquet as pq

        table = pa.table(
            {
                "id": [1],
                "series": pa.array([[1.0, float("nan")]], type=pa.list_(pa.float64())),
                "meta": pa.array([{"a": float("nan")}]),
            }
        )
        path = tmp_path / "both.parquet"
        pq.write_table(table, path)
        columns, before_rows, after_rows, changed = RepairEngine.table_diff(
            path, path, DataSourceType.PARQUET
        )
        assert changed == [], f"identical bytes reported a change: {changed}"
        assert columns == ["id", "series", "meta"], columns
        assert isinstance(before_rows[0][1], list) and math.isnan(before_rows[0][1][1])
        # `==` on the rows themselves would be a lie in the other direction: NaN != NaN, so the
        # assertion must use the same predicate the audit uses.
        assert rows_equal(before_rows[0], after_rows[0])

    def test_a_generated_suffix_never_collides_with_a_real_column_name(
        self, tmp_path: Path
    ) -> None:
        """`a,a,a_1` 里第三个列**本来就叫** `a_1`；给第二个 `a` 造的别名必须避开它。"""
        b = tmp_path / "before.csv"
        a = tmp_path / "after.csv"
        b.write_text("a,a,a_1\n1,2,3\n4,  x ,5\n", encoding="utf-8")
        a.write_text('"a","a_1","a_1_1"\n1,"2",3\n4,"x",5\n', encoding="utf-8")
        columns, before_rows, after_rows, changed = RepairEngine.table_diff(
            b, a, DataSourceType.CSV
        )
        assert columns == ["a", "a_1", "a_1_1"], columns
        assert before_rows == [["1", "2", "3"], ["4", "  x ", "5"]], before_rows
        assert after_rows == [["1", "2", "3"], ["4", "x", "5"]], after_rows
        assert changed == [1], f"a column was read one slot over: {changed}"

    def test_alignment_survives_a_rename_rule_the_reader_does_not_share(
        self, tmp_path: Path
    ) -> None:
        """本机实测：`a_1,a,a` 在连接器侧定成 `a_1`/`a`/`a_1_1`，而读侧的别名规则给的是 `a_2`。

        两套写法不可能都对，所以 CSV 的对齐不许经过名字——列顺序才是这类格式的 schema。
        显示用的列名取自副本（那是连接器自己写下的名字）。
        """
        b = tmp_path / "before.csv"
        a = tmp_path / "after.csv"
        b.write_text("a_1,a,a\n1,2,3\n4,5,  x \n", encoding="utf-8")
        a.write_text('"a_1","a","a_1_1"\n1,"2",3\n4,"5","x"\n', encoding="utf-8")
        columns, before_rows, after_rows, changed = RepairEngine.table_diff(
            b, a, DataSourceType.CSV
        )
        assert columns == ["a_1", "a", "a_1_1"], f"display names must come from the copy: {columns}"
        assert before_rows == [["1", "2", "3"], ["4", "5", "  x "]], before_rows
        assert after_rows == [["1", "2", "3"], ["4", "5", "x"]], after_rows
        assert changed == [1], changed

    def test_a_repair_that_drops_a_column_shows_it_as_dropped(self, tmp_path: Path) -> None:
        """列数不等时按位置补齐：少掉的那一列要显示成 `value -> None`，不能被截掉。"""
        b = tmp_path / "before.csv"
        a = tmp_path / "after.csv"
        b.write_text("id,name,note\n1,x,old\n2,y,chg\n", encoding="utf-8")
        a.write_text("id,name\n1,x\n2,y\n", encoding="utf-8")
        columns, before_rows, after_rows, changed = RepairEngine.table_diff(
            b, a, DataSourceType.CSV
        )
        assert columns == ["id", "name", "note"], columns
        assert before_rows[0] == ["1", "x", "old"], before_rows
        assert after_rows[0] == ["1", "x", None], after_rows
        assert changed == [0, 1], changed

    def test_preview_refuses_an_unrepairable_source_with_the_same_words_as_apply(
        self, tmp_path: Path
    ) -> None:
        """`apply` 的文案承诺"propose and preview work"，而 `preview` 也调 `_suffix`。"""
        import sqlite3

        path = tmp_path / "t.db"
        conn = sqlite3.connect(path)
        conn.execute("create table t (name text)")
        conn.execute("insert into t values ('  a ')")
        conn.commit()
        conn.close()
        handle = default_registry().open(
            DataSourceSpec(source_type=DataSourceType.SQLITE, path=path, table_name="t")
        )
        context = DetectionContext(
            dataset_id="t", table_name="t", columns=handle.schema().column_names, handle=handle
        )
        proposal = RepairProposal(
            proposal_id="prop_test",
            issue_id="iss_test",
            operation=RepairOperation.TRIM_WHITESPACE,
            target_columns=["name"],
            estimated_rows_changed=1,
        )
        with pytest.raises(ValueError, match="does not support sqlite"):
            RepairEngine().preview(proposal, context, DetectorRegistry())

    def test_a_refused_apply_creates_no_directory(self, tmp_path: Path) -> None:
        import sqlite3

        path = tmp_path / "proj" / "t.db"
        path.parent.mkdir(parents=True)
        conn = sqlite3.connect(path)
        conn.execute("create table t (name text)")
        conn.execute("insert into t values ('  a ')")
        conn.commit()
        conn.close()
        import sqlite3

        handle = default_registry().open(
            DataSourceSpec(source_type=DataSourceType.SQLITE, path=path, table_name="t")
        )
        context = DetectionContext(
            dataset_id="t", table_name="t", columns=handle.schema().column_names, handle=handle
        )
        proposal = RepairProposal(
            proposal_id="prop_test",
            issue_id="iss_test",
            operation=RepairOperation.TRIM_WHITESPACE,
            target_columns=["name"],
            estimated_rows_changed=1,
        )
        with pytest.raises(ValueError, match="does not support sqlite"):
            RepairEngine().apply(proposal, context, path.parent)
        assert not (path.parent / ".datasentry").exists(), (
            "the refused repair still created the workspace store directory"
        )


class TestSharedCellPredicate:
    """复核 F9：四面必须用同一个"这格变了吗"。

    顶层 NaN 在产品里到不了端到端（数值列含 NaN 的 Parquet 在扫描阶段就崩，另记候选 F-7），
    所以这条不能只靠一句"两面共用一个谓词"的说法，得各画一面钉住。
    """

    def test_the_workbench_does_not_highlight_a_nan_cell_as_changed(self) -> None:
        from datasentry.ui import _diff_table

        # Two separate NaN objects, as `to_pylist` returns them: one shared object would let a
        # plain `!=` pass on CPython's pointer-equality fast path and hide the bug pinned here.
        html = _diff_table(
            ["id", "series"],
            [[1, [1.0, float("nan")]]],
            [[1, [1.0, float("nan")]]],
            [0],
            lang="en",
        )
        assert "diff-del" not in html and "diff-add" not in html, html

    def test_the_cli_does_not_list_a_nan_cell_as_changed(
        self,
        tmp_path: Path,
        monkeypatch: pytest.MonkeyPatch,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        import argparse

        from datasentry.client import DataSentry

        source = tmp_path / "d.tsv"
        source.write_text("id\tname\n1\t  Alice \n2\tBob\n", encoding="utf-8")
        client = DataSentry(project=tmp_path)
        try:
            _run, _d, _issues = client.scan_file(source)
            issue_id = next(
                i.id
                for i in client.list_issues(scan_run_id=_run.id)
                if "leading_or_trailing_whitespace" in ",".join(i.detector_ids)
            )
            rep = client.repair_apply(issue_id, source)
            _r, columns, before_rows, after_rows, changed = client.repair_diff(rep.id)
            patched = (
                rep,
                [*columns, "series"],
                [[*row, float("nan")] for row in before_rows],
                [[*row, float("nan")] for row in after_rows],
                changed,
            )
            monkeypatch.setattr(client, "repair_diff", lambda _rid: patched)
            monkeypatch.setattr("datasentry.cli.DataSentry", lambda _project: client)
            from datasentry.cli import _cmd_repair_diff

            _cmd_repair_diff(argparse.Namespace(project=tmp_path, run_id=rep.id, format="text"))
        finally:
            client.close()
        printed = capsys.readouterr().out
        assert "name:" in printed, printed
        assert "series:" not in printed, (
            f"the CLI listed an unchanged NaN cell as changed: {printed}"
        )

    def test_one_predicate_answers_which_cells_moved(self) -> None:
        """两面（UI 高亮、CLI 逐格）都调这一个函数；NaN 语义只收在一处。"""
        # distinct NaN objects each side, as `to_pylist` hands them over
        assert changed_cells([1, float("nan"), "  a "], [1, float("nan"), "a"]) == {2}
        assert changed_cells([float("nan")], [float("nan")]) == set()
        assert changed_cells([None], [""]) == {0}
        assert changed_cells([1], [1, 2]) == {1}


class TestCopyKeepsTheSourceShape:
    """复核 B-1 的写侧与 B-5：副本必须还是那种文件，修不了就不许先落盘。

    `_suffix` 只认 CSV/PARQUET/JSONL/XLSX，其余一律回 `.csv`；`_write_table` 的 CSV 分支不带
    任何 delimiter。于是一份 `.tsv` 的"修复副本"是逗号分隔的另一种文件——而
    `ui.repair_copy_note` 对用户承诺的是"写一份副本"。SQLite 更糟：先落两个工件再失败，
    磁盘上留下无人认领的字节，`RepairRun` 一条没存，既不能 diff 也不能回滚。
    """

    def _whitespace_issue_id(self, client: object, run_id: str) -> str:
        issues = client.list_issues(scan_run_id=run_id)
        return next(
            i.id for i in issues if "leading_or_trailing_whitespace" in ",".join(i.detector_ids)
        )

    def test_a_tsv_repair_produces_a_tsv_copy(self, tmp_path: Path) -> None:
        from datasentry.client import DataSentry

        src = tmp_path / "dirty.tsv"
        src.write_text("id\tname\n1\t  Alice \n2\tBob\n", encoding="utf-8")
        client = DataSentry(project=tmp_path)
        try:
            run, _ds, _issues = client.scan_file(src)
            rep = client.repair_apply(self._whitespace_issue_id(client, run.id), src)
            copy_path = Path(rep.rollback_artifact).with_name(
                Path(rep.rollback_artifact).name.replace(".before", "")
            )
            assert copy_path.suffix == ".tsv", f"artefact renamed to {copy_path.suffix}"
            header = copy_path.read_text(encoding="utf-8").splitlines()[0]
            # What has to survive is the *shape*: same delimiter, same columns. pyarrow quotes a
            # header field when the delimiter is a tab, which is the same data either way.
            bare = header.replace('"', "")
            assert "\t" in bare and "," not in bare, f"dialect changed: {header!r}"
            assert [c.strip('"') for c in header.split("\t")] == ["id", "name"], header
            _scan, report = client.repair_verify(rep.id)
            assert report["verify_issue_count"] < report["source_issue_count"], (
                "the copy no longer verifies as a repaired version of the same shape"
            )
        finally:
            client.close()

    def test_apply_refuses_an_unrepairable_source_before_writing_anything(
        self, tmp_path: Path
    ) -> None:
        import sqlite3

        from datasentry.client import DataSentry

        src = tmp_path / "orders.db"
        conn = sqlite3.connect(src)
        conn.execute("CREATE TABLE orders (id INTEGER, name TEXT)")
        conn.execute("INSERT INTO orders VALUES (1, '  Alice ')")
        conn.commit()
        conn.close()
        client = DataSentry(project=tmp_path)
        try:
            run, _ds, _issues = client.scan_file(src, table_name="orders")
            issue_id = self._whitespace_issue_id(client, run.id)
            repairs = tmp_path / ".datasentry" / "repairs"
            before = sorted(p.name for p in repairs.glob("*")) if repairs.exists() else []
            with pytest.raises(ValueError, match="repair does not support"):
                client.repair_apply(issue_id, src)
            after = sorted(p.name for p in repairs.glob("*")) if repairs.exists() else []
            assert after == before, f"artefacts written for a refused repair: {after}"
            assert client.list_repair_runs() == [], "a refused repair still persisted a run"
        finally:
            client.close()


class TestDiffDoesNotInventChanges:
    """D2-10：diff 会报出没发生的变更——两个方向各一次。

    `D2-06` 的第一症状是"该报的没报"（假阴性，r7 已修）；这两例是"不该报的报了"。
    P31-A 之后这份 diff 出现在每次 apply 的响应页上，所以假阳性的暴露面比原来更大。
    """

    def test_two_identical_nan_bearing_files_report_no_change(self, tmp_path: Path) -> None:
        """`float('nan') != float('nan')` 恒真，所以 `!=` 比较会在两份逐字节相同的工件上报变更。"""
        import math

        import pyarrow as pa
        import pyarrow.parquet as pq

        table = pa.table({"v": [1.0, float("nan"), 3.0]})
        b = tmp_path / "before.parquet"
        a = tmp_path / "after.parquet"
        pq.write_table(table, b)
        pq.write_table(table, a)
        assert b.read_bytes() == a.read_bytes()
        columns, before_rows, after_rows, changed = RepairEngine.table_diff(
            b, a, DataSourceType.PARQUET
        )
        assert columns == ["v"]
        assert changed == [], f"NaN claimed a change that did not happen: {before_rows}"
        assert math.isnan(before_rows[1][0])
        assert math.isnan(after_rows[1][0])

    def test_a_real_repair_does_not_report_the_nan_row_as_changed(self, tmp_path: Path) -> None:
        """D2-10(b) 的端到端形态——G-2 修好之后这条才可能在产品里跑，而不只在工件对上证明。

        修复只动 `name` 的两格；含 NaN 的第 2 行一个字都没变，而它正是过去 `!=` 会报变更的那一行。
        此前台账只能声明"单元层证明"，因为含 NaN 的数值列在扫描阶段就整体失败。
        """
        import pyarrow as pa
        import pyarrow.parquet as pq

        from datasentry.client import DataSentry

        src = tmp_path / "dirty.parquet"
        pq.write_table(
            pa.table(
                {
                    "name": ["  Alice ", "Bob", "  Carol "],
                    "qty": pa.array([1.0, float("nan"), 3.0], type=pa.float64()),
                }
            ),
            src,
        )
        client = DataSentry(project=tmp_path)
        try:
            run, _dataset, _issues = client.scan_file(src)
            assert run.status == "completed", f"a NaN column failed the scan: {run.error}"
            issue_id = next(
                i.id
                for i in client.list_issues(scan_run_id=run.id)
                if "leading_or_trailing_whitespace" in ",".join(i.detector_ids)
            )
            rep = client.repair_apply(issue_id, src)
            _r, _columns, before_rows, after_rows, changed = client.repair_diff(rep.id)
        finally:
            client.close()
        assert changed == [0, 2], changed
        # `==` would lie again here (NaN != NaN across two reads), so the assertion uses the very
        # predicate the audit is built on.
        assert rows_equal(before_rows[1], after_rows[1]), (
            "the untouched NaN row was claimed as rewritten"
        )

    def test_an_interior_nulled_xlsx_row_stays_on_its_own_line(self, tmp_path: Path) -> None:
        """读取端把**中间**的整空行丢掉，等于给 after 侧重排了行号，按位比较就凭空造出改写。"""
        from openpyxl import Workbook

        def workbook(name: str, values: list[object]) -> Path:
            wb = Workbook()
            sheet = wb.active
            sheet.title = "Data"
            sheet.append(["date"])
            for value in values:
                sheet.append([value])
            path = tmp_path / name
            wb.save(path)
            return path

        b = workbook("before.xlsx", ["2024-01-01", "not-a-date", "2024-02-02"])
        a = workbook("after.xlsx", ["2024-01-01", None, "2024-02-02"])
        columns, before_rows, after_rows, changed = RepairEngine.table_diff(
            b, a, DataSourceType.XLSX
        )
        assert len(after_rows) == 3, f"a row vanished from the snapshot's counterpart: {after_rows}"
        assert columns == ["date"]
        assert changed == [1], (
            f"misaligned rows fabricated {changed}: {before_rows} vs {after_rows}"
        )
        assert before_rows[1] == ["not-a-date"] and after_rows[1] == [None]


class TestCopyIsTheSameKindOfFile:
    """副本必须还能被同一个连接器打开——`_write_table` 的分支一旦漏了 return，
    后面那段 CSV 写入就会把刚存好的 xlsx 覆盖成 CSV 字节，`apply` 随后指纹化副本时炸。

    这条是我自己在这批改动里写出来的形状错误：单测只喂两份工件给 `table_diff`，看不见
    apply 之后"副本是不是还能打开"。
    """

    def _apply(self, tmp_path: Path, source: Path) -> object:
        from datasentry.client import DataSentry

        self._ws = tmp_path
        client = DataSentry(project=tmp_path)
        try:
            run, _ds, _issues = client.scan_file(source)
            issues = client.list_issues(scan_run_id=run.id)
            issue_id = next(
                i.id for i in issues if "leading_or_trailing_whitespace" in ",".join(i.detector_ids)
            )
            return client.repair_apply(issue_id, source)
        finally:
            client.close()

    def _copy_of(self, rep: object) -> Path:
        artifact = Path(rep.rollback_artifact)
        return artifact.with_name(artifact.name.replace(".before", ""))

    def test_an_xlsx_repair_leaves_an_openable_xlsx_copy(self, tmp_path: Path) -> None:
        from openpyxl import Workbook

        source = tmp_path / "orders.xlsx"
        wb = Workbook()
        sheet = wb.active
        sheet.append(["id", "name"])
        sheet.append([1, "  Alice "])
        sheet.append([2, "Bob"])
        wb.save(source)

        rep = self._apply(tmp_path, source)
        copy_path = self._copy_of(rep)
        assert copy_path.suffix == ".xlsx"
        assert copy_path.read_bytes()[:2] == b"PK", "the copy is no longer a zip container"
        from openpyxl import load_workbook

        rows = list(load_workbook(copy_path).worksheets[0].iter_rows(values_only=True))
        assert rows[0] == ("id", "name")
        assert rows[1] == (1, "Alice"), f"the copy lost the repair: {rows}"

    def test_a_parquet_repair_leaves_an_openable_parquet_copy(self, tmp_path: Path) -> None:
        import pyarrow as pa
        import pyarrow.parquet as pq

        source = tmp_path / "orders.parquet"
        pq.write_table(pa.table({"id": [1, 2], "name": ["  Alice ", "Bob"]}), source)
        rep = self._apply(tmp_path, source)
        copy_path = self._copy_of(rep)
        assert copy_path.suffix == ".parquet"
        assert copy_path.read_bytes()[:4] == b"PAR1", "the copy is no longer Parquet"
        table = pq.read_table(copy_path)
        assert table.column("name").to_pylist() == ["Alice", "Bob"]

    def test_a_jsonl_repair_is_refused_until_carry_supports_it(self, tmp_path: Path) -> None:
        """改钉（2026-10-02）：JSONL 的副本曾由 DuckDB 投影整表重写，而复核实测它会改动
        提案没碰的键（ISO 时间的 `T`、缺失键补 null）。在逐格搬运支持 JSONL 之前，
        正确行为是拒绝，而不是写一份 unauthorized 的副本。"""
        import json

        source = tmp_path / "orders.jsonl"
        source.write_text(
            "\n".join(
                [
                    json.dumps({"id": 1, "name": "  Alice "}, ensure_ascii=False),
                    json.dumps({"id": 2, "name": "Bob"}, ensure_ascii=False),
                ]
            )
            + "\n",
            encoding="utf-8",
        )
        with pytest.raises(ValueError, match="JSONL"):
            self._apply(tmp_path, source)

    def _verify(self, rep: object) -> tuple:
        from datasentry.client import DataSentry

        client = DataSentry(project=self._ws)
        try:
            return client.repair_verify(rep.id)
        finally:
            client.close()


class TestCopyChangesOnlyTheAuthorisedCells:
    """候选 `G-1`：一次 TRIM 之外，副本还静默重写了**非目标列**的格。

    本机实测（`/tmp/g1/measure.py`）：CSV/TSV 的 `qty` 被写成 `1e5->100000`、`2.50e3->2500`、
    `1.50->1.5`，`flag` 被写成 `TRUE->true`；混合列的 XLSX 把整表数字写成文本（`1->'1'`、
    `10->'10'`）。更关键的是**提案没碰的那一行也被改写**——CSV 用例里 `changed=[0,1,2]` 而
    `operations` 只有两格。JSONL 与 Parquet 实测无附带改写（本行同时钉住它们不回归）。

    规则：副本 = 源文件自己的格子 + 只被授权的目标列替换。所以
    `changed ⊆ 提案覆盖的行`、且 `changed` 里每一格的差异都必须能在 `operations` 里找到。
    """

    def _apply(self, tmp_path: Path, source: Path) -> tuple[object, list[str], list, list, list]:
        from datasentry.client import DataSentry

        client = DataSentry(project=tmp_path)
        try:
            run, _dataset, _issues = client.scan_file(source)
            assert run.status == "completed", run.error
            issue_id = next(
                i.id
                for i in client.list_issues(scan_run_id=run.id)
                if "leading_or_trailing_whitespace" in ",".join(i.detector_ids)
            )
            rep = client.repair_apply(issue_id, source)
            _r, columns, before_rows, after_rows, changed = client.repair_diff(rep.id)
            return rep, columns, before_rows, after_rows, changed
        finally:
            client.close()

    def test_a_csv_repair_leaves_the_other_columns_as_written(self, tmp_path: Path) -> None:
        source = tmp_path / "d.csv"
        source.write_text(
            "name,qty,flag,when\n"
            "  Alice ,1e5,TRUE,2024-1-5\n"
            "Bob,2.50e3,true,2024-01-06\n"
            "  Carol ,1.50,FALSE,2024-02-30\n",
            encoding="utf-8",
        )
        rep, columns, before_rows, after_rows, changed = self._apply(tmp_path, source)
        assert columns == ["name", "qty", "flag", "when"], columns
        assert changed == [0, 2], (
            f"a two-cell proposal rewrote {changed}: {before_rows} -> {after_rows}"
        )
        assert before_rows[1] == after_rows[1], "an untouched row was rewritten"
        for i in changed:
            for j, column in enumerate(columns):
                if column == "name":
                    continue
                assert before_rows[i][j] == after_rows[i][j], (
                    f"non-target column {column} rewritten: {before_rows[i]} -> {after_rows[i]}"
                )
        assert [(o.column, o.before, o.after) for o in rep.operations] == [
            ("name", "  Alice ", "Alice"),
            ("name", "  Carol ", "Carol"),
        ], rep.operations

    def test_a_tsv_repair_keeps_semicolon_and_scientific_text(self, tmp_path: Path) -> None:
        source = tmp_path / "d.tsv"
        source.write_text(
            "name\tqty\tflag\n  Alice \t1e5\tTRUE\nBob\t2.50e3\ttrue\n  Carol \t1.50\tFALSE\n",
            encoding="utf-8",
        )
        _rep, columns, before_rows, after_rows, changed = self._apply(tmp_path, source)
        assert columns == ["name", "qty", "flag"], columns
        assert changed == [0, 2], changed
        assert before_rows[1] == after_rows[1], "an untouched row was rewritten"
        assert [row[1] for row in after_rows] == ["1e5", "2.50e3", "1.50"], after_rows
        assert [row[2] for row in after_rows] == ["TRUE", "true", "FALSE"], after_rows

    def test_a_mixed_column_xlsx_repair_keeps_cell_types(self, tmp_path: Path) -> None:
        from openpyxl import Workbook

        source = tmp_path / "m.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.append(["name", "mixed", "qty"])
        sheet.append(["  Alice ", 1, 10])
        sheet.append(["Bob", "N/A", 20])
        sheet.append(["  Carol ", 3.5, 30])
        book.save(source)
        _rep, columns, before_rows, after_rows, changed = self._apply(tmp_path, source)
        assert columns == ["name", "mixed", "qty"], columns
        assert changed == [0, 2], f"numbers were turned into text: {before_rows} -> {after_rows}"
        assert after_rows[1] == ["Bob", "N/A", 20], after_rows[1]
        assert after_rows[0][1:] == [1, 10], after_rows[0]
        assert isinstance(after_rows[0][1], int), "a numeric cell became text"

    def test_the_copy_still_repairs_the_target_column(self, tmp_path: Path) -> None:
        """不回归的反面：授权的那一格必须真的被修好。"""
        source = tmp_path / "keep.csv"
        source.write_text("name,qty\n  Alice ,1e5\nBob,2\n", encoding="utf-8")
        _rep, columns, before_rows, after_rows, changed = self._apply(tmp_path, source)
        assert columns == ["name", "qty"], columns
        assert changed == [0], changed
        assert before_rows[0] == ["  Alice ", "1e5"], before_rows[0]
        assert after_rows[0] == ["Alice", "1e5"], after_rows[0]

    def test_the_jsonl_repair_refuses_rather_than_rewrite(self, tmp_path: Path) -> None:
        """改钉（2026-10-02）：同键集、无 ISO-T 字符串的记录确实测不到附带改写，但那只是
        DuckDB 投影恰好保值的特例——异键集与时间戳都会被改写而 `operations` 不记
        （`REVIEW_CARRY_G1_R1.md` F2）。逐格搬运支持 JSONL 之前一律拒绝。"""
        import json

        source = tmp_path / "d.jsonl"
        rows = [
            {"name": "  Alice ", "qty": 100000.0, "flag": True},
            {"name": "Bob", "qty": 2500.0, "flag": False},
            {"name": "  Carol ", "qty": 1.5, "flag": True},
        ]
        source.write_text("\n".join(json.dumps(r) for r in rows) + "\n", encoding="utf-8")
        with pytest.raises(ValueError, match="JSONL"):
            self._apply(tmp_path, source)

    def test_a_parquet_copy_gains_no_extra_changes(self, tmp_path: Path) -> None:
        import pyarrow as pa
        import pyarrow.parquet as pq

        source = tmp_path / "d.parquet"
        pq.write_table(
            pa.table(
                {
                    "name": ["  Alice ", "Bob", "  Carol "],
                    "qty": pa.array([100000.0, 2500.0, 1.5], type=pa.float64()),
                    "flag": [True, False, True],
                }
            ),
            source,
        )
        _rep, _columns, _before, after, changed = self._apply(tmp_path, source)
        assert changed == [0, 2], changed
        assert [row[1] for row in after] == [100000.0, 2500.0, 1.5], after

    def test_a_source_whose_rows_cannot_be_read_verbatim_is_refused_before_writing(
        self, tmp_path: Path
    ) -> None:
        """合并的前提是"逐格搬源文件"。读不出对齐的行，就宁可不写，也不 produce 一份错位的副本。"""
        from datasentry.client import DataSentry

        source = tmp_path / "ragged.csv"
        source.write_text("id,name\n1,Alice,extra\n2,  Bob ,x\n", encoding="utf-8")
        client = DataSentry(project=tmp_path)
        try:
            run, _dataset, _issues = client.scan_file(source)
            issues = [
                i
                for i in client.list_issues(scan_run_id=run.id)
                if "leading_or_trailing_whitespace" in ",".join(i.detector_ids)
            ]
            with pytest.raises(ValueError, match="cell-by-cell") as exc:
                client.repair_apply(issues[0].id, source)
            assert not list((tmp_path / ".datasentry" / "repairs").glob("*.csv")), (
                "a misaligned copy was written anyway"
            )
            assert exc.value is not None
        finally:
            client.close()


class TestCarryAlignsWithTheProjection:
    """逐格搬运的前提：搬的必须是与投影**同一批行**。

    画像读取器（`_read_artefact`）按审计口径裁掉尾部空行，而连接器的视图按表的存储范围保留它们
    ——本机实测一份尾部有一空行的 XLSX：读侧 2 行、投影 3 行。照读侧搬就会**错位**，
    更糟的是我会把一份本来可修的文件拒掉。所以副本读取自己的行，不借审计的那把尺子。
    """

    def test_an_xlsx_with_a_trailing_blank_row_still_repairs(self, tmp_path: Path) -> None:
        from openpyxl import Workbook

        source = tmp_path / "t.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.append(["name", "qty"])
        sheet.append(["  Alice ", 1])
        sheet.append(["Bob", 2])
        sheet.append([None, None])
        book.save(source)

        from datasentry.client import DataSentry

        client = DataSentry(project=tmp_path)
        try:
            run, _dataset, _issues = client.scan_file(source)
            issue_id = next(
                i.id
                for i in client.list_issues(scan_run_id=run.id)
                if "leading_or_trailing_whitespace" in ",".join(i.detector_ids)
            )
            rep = client.repair_apply(issue_id, source)
            _r, columns, before_rows, after_rows, changed = client.repair_diff(rep.id)
        finally:
            client.close()
        assert columns == ["name", "qty"], columns
        assert changed == [0], f"rows were carried out of alignment: {before_rows} -> {after_rows}"
        assert after_rows[0] == ["Alice", 1], after_rows[0]
        assert after_rows[1] == ["Bob", 2], after_rows[1]

    def test_the_guard_refuses_a_projection_that_does_not_match_the_source_rows(
        self, tmp_path: Path
    ) -> None:
        """守卫本身要可证伪：直接喂一个行数对不上的投影，它必须拒绝而不是错位搬。"""
        import pyarrow as pa

        from datasentry_core.connectors import DataSourceSpec, default_registry
        from datasentry_core.detectors import DetectionContext

        source = tmp_path / "g.csv"
        source.write_text("name,qty\n  Alice ,1e5\nBob,2\n", encoding="utf-8")
        handle = default_registry().open(
            DataSourceSpec(source_type=DataSourceType.CSV, path=source)
        )
        try:
            context = DetectionContext(
                dataset_id="g", table_name=None, columns=["name", "qty"], handle=handle
            )
            proposal = RepairProposal(
                proposal_id="prop_g",
                issue_id="iss_g",
                operation=RepairOperation.TRIM_WHITESPACE,
                target_columns=["name"],
                estimated_rows_changed=1,
            )
            short = pa.Table.from_pydict({"name": ["Alice"], "qty": ["1e5"]})
            with pytest.raises(ValueError, match="row"):
                RepairEngine._copy_cells(context, short, proposal, source)
        finally:
            handle.close()


class TestCarryRefusesWhatItCannotCarryFaithfully:
    """对并行会话 G-1 实施的复核发现（R1-F1/F2/F4）的处置：搬不动就拒绝，不许错着搬。

    实测三例：① XLSX 重名表头（连接器 `dict(zip)` 塌缩）让逐格搬运把修复结果写进**错误的列**——
    复制体第三格表头变成 `None`、第二格数据被改而 `operations` 声称的 before/after 与实际被改的格
    对不上；② CSV 重名表头同理；③ JSONL 的"豁免"依据不实——非目标键被改写（ISO 时间的 `T` 变
    空格、缺失键被补 `"key": null`）。逐格搬运的承诺是"没被授权的格一个都不动"；凡是对齐做不到
    的输入，正确行为是拒绝修复，而不是写一份 unauthorized 的副本。
    """

    def _apply_expect_refusal(self, tmp_path: Path, name: str, write: object) -> None:
        from datasentry.client import DataSentry

        src = write(tmp_path / name)
        client = DataSentry(project=tmp_path)
        try:
            run_, _d, _i = client.scan_file(src)
            assert run_.status == "completed", run_.error
            issue_id = next(
                i.id
                for i in client.list_issues(scan_run_id=run_.id)
                if "leading_or_trailing_whitespace" in ",".join(i.detector_ids)
            )
            with pytest.raises(ValueError, match="carry"):
                client.repair_apply(issue_id, src)
            repairs = tmp_path / ".datasentry" / "repairs"
            assert (
                not any(p.name.startswith("rep_") for p in repairs.glob("*"))
                if repairs.exists()
                else True
            ), "the refused repair still wrote artefacts"
        finally:
            client.close()

    def test_an_xlsx_sheet_with_duplicate_headers_is_refused(self, tmp_path: Path) -> None:
        def write(p: Path) -> Path:
            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.append(["a_1", "a", "a"])
            ws.append(["1", "  x ", "  y "])
            wb.save(p)
            return p

        self._apply_expect_refusal(tmp_path, "dup.xlsx", write)

    def test_a_csv_with_duplicate_headers_is_refused(self, tmp_path: Path) -> None:
        def write(p: Path) -> Path:
            p.write_text("a,a,a_1\n1,  x ,3\n4,5,  y \n", encoding="utf-8")
            return p

        self._apply_expect_refusal(tmp_path, "dup.csv", write)

    def test_a_jsonl_source_is_refused_until_carry_supports_it(self, tmp_path: Path) -> None:
        def write(p: Path) -> Path:
            p.write_text(
                '{"name":"  Alice ","ts":"2024-01-05T08:00:00"}\n{"name":"Bob"}\n',
                encoding="utf-8",
            )
            return p

        self._apply_expect_refusal(tmp_path, "d.jsonl", write)
